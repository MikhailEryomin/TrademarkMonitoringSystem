import pandas as pd
import json
import os
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT_FILE = "dataset_for_customer.xlsx"

def format_contacts(contacts: dict) -> str:
    """Преобразует словарь контактов в читаемую строку."""
    parts = []
    
    if contacts.get('phones'):
        # Убираем дубликаты и соединяем
        phones = ", ".join(list(set(contacts['phones'])))
        parts.append(f"Тел: {phones}")
        
    if contacts.get('emails'):
        emails = ", ".join(list(set(contacts['emails'])))
        parts.append(f"Email: {emails}")
        
    if contacts.get('inn'):
        inn = ", ".join(list(set(contacts['inn'])))
        parts.append(f"ИНН: {inn}")
        
    return "\n".join(parts) if parts else "Отсутствуют"

def generate_comment(status: str, title: str) -> str:
    """Генерирует твой предварительный комментарий."""
    if status == 'parked':
        return "Парковочная страница / Заглушка"
    if status == 'redirect':
        return "Переадресация на сторонний ресурс"
    if "официальный" in title.lower():
        return "Заявляют об официальности"
    return "Требует проверки"

def json_to_dataframe(json_path: str, target_brand_name: str, start_id: int = 1) -> pd.DataFrame:
    """Читает JSON и готовит DataFrame в нужной структуре."""
    
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    rows = []
    current_id = start_id
    
    for item in data:
        # Пропускаем явный мусор, если нужно (или оставляем всё)
        # if item.get('status') == 'dead': continue

        row = {
            "ID": current_id,
            "ТЗ": target_brand_name,
            "URL": item.get('url'),
            "Title": item.get('title', ''),
            "Контакты (почта, тел. номера, ИНН)": format_contacts(item.get('contacts', {})),
            "Комментарии (Ваши)": generate_comment(item.get('status', ''), item.get('title', '')),
            "Label (Выберите из списка)": "",  # Пусто для заказчика
            "Комментарии заказчика": "",       # Пусто для заказчика
            "Фрагмент контента": item.get('content_sample', '')[:300] + "..." # Обрезаем, чтобы не раздувать Excel
        }
        rows.append(row)
        current_id += 1
        
    return pd.DataFrame(rows)

def append_to_excel(df: pd.DataFrame, filename: str):
    """
    Добавляет данные в Excel. 
    Если файла нет - создает красивый с выпадающими списками.
    Если есть - просто дописывает вниз.
    """
    
    # 1. Если файла нет, создаем его с нуля с форматированием
    if not os.path.exists(filename):
        # Используем xlsxwriter для начального красивого создания
        with pd.ExcelWriter(filename, engine='xlsxwriter') as writer:
            df.to_excel(writer, sheet_name='Dataset', index=False)
            workbook = writer.book
            worksheet = writer.sheets['Dataset']
            
            # Форматы
            header_format = workbook.add_format({
                'bold': True, 'text_wrap': True, 'valign': 'top', 
                'fg_color': '#D9E1F2', 'border': 1
            })
            text_wrap_format = workbook.add_format({'text_wrap': True, 'valign': 'top'})
            
            # Применяем форматы и ширину колонок
            worksheet.set_column('A:A', 5, text_wrap_format)   # ID
            worksheet.set_column('B:B', 15, text_wrap_format)  # ТЗ
            worksheet.set_column('C:C', 30, text_wrap_format)  # URL
            worksheet.set_column('D:D', 25, text_wrap_format)  # Title
            worksheet.set_column('E:E', 30, text_wrap_format)  # Контакты
            worksheet.set_column('F:F', 20, text_wrap_format)  # Твои комменты
            worksheet.set_column('G:G', 15, text_wrap_format)  # Label
            worksheet.set_column('H:H', 25, text_wrap_format)  # Комменты заказчика
            worksheet.set_column('I:I', 50, text_wrap_format)  # Контент
            
            # Записываем заголовки с форматом
            for col_num, value in enumerate(df.columns.values):
                worksheet.write(0, col_num, value, header_format)
                
            # Добавляем выпадающий список (Data Validation) для колонки Label (G)
            # Примечание: xlsxwriter делает это немного иначе, чем openpyxl, 
            # но проще сделать это пост-фактум через openpyxl ниже.
            
        print(f"[+] Created new file: {filename}")

        # Открываем через openpyxl, чтобы добавить выпадающий список (это надежнее)
        wb = load_workbook(filename)
        ws = wb['Dataset']
        
        # Создаем валидацию (0, 1, 2, 3)
        dv = DataValidation(type="list", formula1='"0-Safe,1-Parked,2-Phishing,3-Grey"', allow_blank=True)
        dv.error = 'Выберите значение из списка'
        dv.errorTitle = 'Некорректный ввод'
        
        # Применяем к колонке G (Label) от 2 строки до 1000
        dv.add('G2:G1000')
        ws.add_data_validation(dv)
        wb.save(filename)

    # 2. Если файл есть, добавляем данные (Append)
    else:
        # Читаем существующий, чтобы узнать последний ID
        existing_df = pd.read_excel(filename)
        last_id = existing_df['ID'].max() if not existing_df.empty else 0
        
        # Обновляем ID новых строк
        df['ID'] = range(last_id + 1, last_id + 1 + len(df))
        
        # Дописываем (mode='a' - append, if_sheet_exists='overlay' - не затирать)
        with pd.ExcelWriter(filename, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
            # Находим последнюю строку
            start_row = writer.sheets['Dataset'].max_row
            df.to_excel(writer, sheet_name='Dataset', index=False, header=False, startrow=start_row)
            
        print(f"[+] Appended {len(df)} rows to {filename}")

if __name__ == "__main__":
    # # 1. Загружаем Samsung
    # if os.path.exists("output/json/data_samsung.json"):
    #     df_samsung = json_to_dataframe("output/json/data_samsung.json", "Samsung")
    #     append_to_excel(df_samsung, OUTPUT_FILE)
        
    # # 2. Загружаем Sberbank
    # if os.path.exists("output/json/data_sber.json"):
    #     df_sber = json_to_dataframe("output/json/data_sber.json", "Sberbank")
    #     append_to_excel(df_sber, OUTPUT_FILE)

    # # 3. Загружаем Adidas
    # if os.path.exists("output/json/data_adidas.json"):
    #     df_sber = json_to_dataframe("output/json/data_adidas.json", "Adidas")
    #     append_to_excel(df_sber, OUTPUT_FILE)

    # # 3. Загружаем Avito
    # if os.path.exists("output/json/data_avito.json"):
    #     df_sber = json_to_dataframe("output/json/data_avito.json", "Avito")
    #     append_to_excel(df_sber, OUTPUT_FILE)

    # 4. Загружаем Grey
    if os.path.exists("output/json/data_Ozon.json"):
        df_sber = json_to_dataframe("output/json/data_Ozon.json", "Ozon")
        append_to_excel(df_sber, OUTPUT_FILE)